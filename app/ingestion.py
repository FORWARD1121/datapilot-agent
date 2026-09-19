"""Bounded CSV/XLSX parsing without formula evaluation or silent imputation."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from app.core import AppError, Settings


@dataclass
class ParsedTable:
    frame: pd.DataFrame
    audit: dict


def _headers(values: list, settings: Settings) -> list[str]:
    names = [str(value).strip() if value is not None else "" for value in values]
    if not 2 <= len(names) <= settings.max_columns:
        raise AppError("invalid_columns", "A table must have between 2 and the configured maximum columns")
    if any(not x or len(x) > 100 or any(ord(c) < 32 for c in x) for x in names):
        raise AppError("invalid_header", "Column names must be nonempty, bounded, printable text")
    if len(set(names)) != len(names):
        raise AppError("duplicate_header", "Duplicate column names are not allowed")
    return names


def _rows(iterator, settings: Settings) -> tuple[list[str], list[list]]:
    try:
        names = _headers(list(next(iterator)), settings)
    except StopIteration:
        raise AppError("empty_file", "The file contains no table") from None
    rows = []
    for row in iterator:
        if len(row) != len(names):
            raise AppError("ragged_table", "Every record must have the same number of columns")
        rows.append(list(row))
        if len(rows) > settings.max_rows or len(rows) * len(names) > settings.max_cells:
            raise AppError("table_too_large", "Row or cell limit exceeded", 413)
    return names, rows


def _csv(raw: bytes, settings: Settings) -> tuple[list[str], list[list], str]:
    text = None
    encoding = ""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None or "\x00" in text:
        raise AppError("invalid_encoding", "Expected UTF-8 or GB18030 CSV without NUL bytes")
    try:
        names, rows = _rows(csv.reader(io.StringIO(text), strict=True), settings)
    except csv.Error:
        raise AppError("invalid_csv", "Malformed CSV") from None
    return names, rows, encoding


def _xlsx(raw: bytes, settings: Settings) -> tuple[list[str], list[list]]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(x.file_size for x in entries) > settings.max_uncompressed_bytes:
                raise AppError("unsafe_xlsx", "Excel archive exceeds decompression limits", 413)
            for item in entries:
                name = item.filename.lower()
                if (item.flag_bits & 1 or "vbaproject" in name or "externallinks/" in name
                        or item.file_size / max(item.compress_size, 1) > 200):
                    raise AppError("unsafe_xlsx", "Encrypted, external-linked, macro or highly compressed Excel is not accepted")
                if name.endswith((".xml", ".rels")):
                    content = archive.read(item)
                    if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                        raise AppError("unsafe_xlsx", "XML declarations with external entities are not accepted")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise AppError("multiple_sheets", "Export one worksheet per upload")
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > settings.max_columns:
                raise AppError("table_too_large", "Excel column limit exceeded", 413)
            if sheet.max_row and sheet.max_row > settings.max_rows + 1:
                raise AppError("table_too_large", "Excel row limit exceeded", 413)

            def values():
                for row in sheet.iter_rows():
                    if any(cell.data_type in {"f", "e"} for cell in row):
                        raise AppError("excel_formula", "Export values only: formulas and error cells are not accepted")
                    yield [cell.value for cell in row]

            return _rows(iter(values()), settings)
        finally:
            book.close()
    except AppError:
        raise
    except Exception:
        raise AppError("invalid_xlsx", "The file is not a supported Excel workbook") from None


def normalize(frame: pd.DataFrame, drop_duplicates: bool = False) -> ParsedTable:
    """Lossless type inference where possible; no imputation; audit every removal."""
    original_rows = len(frame)
    for name in frame:
        frame[name] = frame[name].map(lambda x: x.strip() if isinstance(x, str) else x)
    frame = frame.replace("", None).dropna(how="all").copy()
    empty_rows = original_rows - len(frame)
    if frame.empty:
        raise AppError("empty_table", "The table contains no nonempty records")
    nonfinite = {}
    inferred = {}
    for name in frame:
        series = frame[name]
        nonempty = series.dropna()
        if nonempty.empty:
            inferred[name] = "empty"
            continue
        text = nonempty.astype(str)
        identifier = name.lower() == "id" or name.lower().endswith("_id")
        leading_zeros = text.str.match(r"^[+-]?0\d+$").any()
        date_hint = bool(re.search(r"date|time|日期|时间", name, re.I))
        iso_dates = text.str.match(r"^\d{4}-\d{2}-\d{2}(?:[ T].*)?$").all()
        if date_hint or iso_dates or pd.api.types.is_datetime64_any_dtype(series):
            converted = pd.to_datetime(series, errors="coerce", format="mixed", utc=True)
            if converted.notna().sum() == len(nonempty):
                frame[name] = converted.dt.tz_convert(None)
                inferred[name] = "datetime"
                continue
        numeric = pd.to_numeric(series, errors="coerce")
        if not identifier and not leading_zeros and numeric.notna().sum() == len(nonempty):
            bad = np.isinf(numeric.to_numpy(dtype=float, na_value=np.nan))
            if bad.any():
                nonfinite[name] = int(bad.sum())
                numeric = numeric.replace([np.inf, -np.inf], np.nan)
            frame[name] = numeric
            inferred[name] = "numeric"
        else:
            frame[name] = series.map(lambda x: None if pd.isna(x) else str(x))
            inferred[name] = "categorical"
    duplicates = int(frame.duplicated().sum())
    if drop_duplicates:
        frame = frame.drop_duplicates()
    return ParsedTable(frame.reset_index(drop=True), {
        "original_rows": original_rows, "empty_rows_removed": empty_rows,
        "duplicate_rows_detected": duplicates,
        "duplicate_rows_removed": duplicates if drop_duplicates else 0,
        "nonfinite_to_null": nonfinite, "inferred_types": inferred,
        "imputation": "none", "whitespace": "trimmed", "timezone": "UTC-naive",
    })


def parse_upload(filename: str, raw: bytes, settings: Settings, drop_duplicates: bool = False) -> ParsedTable:
    if (not filename or len(filename) > 128 or "/" in filename or "\\" in filename
            or any(ord(x) < 32 for x in filename)):
        raise AppError("invalid_filename", "Provide a simple filename without path components")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise AppError("unsupported_format", "Only .csv and .xlsx are supported", 415)
    if not raw:
        raise AppError("empty_file", "The uploaded file is empty")
    if len(raw) > settings.max_upload_bytes:
        raise AppError("file_too_large", "Upload limit exceeded", 413)
    encoding = None
    if suffix == ".csv":
        names, rows, encoding = _csv(raw, settings)
    else:
        names, rows = _xlsx(raw, settings)
    result = normalize(pd.DataFrame(rows, columns=names), drop_duplicates)
    result.audit.update({"format": suffix[1:], "encoding": encoding})
    return result
